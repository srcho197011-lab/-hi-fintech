# -*- coding: utf-8 -*-
"""부록 워크북 검증 — (1) 수식 오류·빈 값 (2) 기대값(원천 예산양식) 대조 (3) 금지 표현 (4) 인쇄 설정
사용: py check_appendix.py <부록.xlsx>   (엑셀 재계산 저장 뒤 · <부록.xlsx>.expect.json 필요) — 종료코드 0이면 통과"""
import io
import json
import re
import sys

from openpyxl import load_workbook

X = sys.argv[1]
exp = json.load(io.open(X + ".expect.json", encoding="utf-8"))
wf = load_workbook(X)
wv = load_workbook(X, data_only=True)
ok = True

# (1) 수식 오류 · 계산값 없음
nform = nerr = nnone = 0
for ws in wf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                nform += 1
                v = wv[ws.title][c.coordinate].value
                if v is None:
                    nnone += 1
                elif isinstance(v, str) and v.startswith("#"):
                    nerr += 1
print(f"수식 {nform}개 · 오류 {nerr} · 값 없음 {nnone}")
ok &= nerr == 0 and nnone == 0

# (2) 기대값 대조(원 단위 · 비율은 1e-6)
worst, bad = 0.0, []
for ref, e in exp.items():
    sh, cell = ref.split("!")
    v = wv[sh][cell].value
    if isinstance(e, (int, float)):
        if not isinstance(v, (int, float)):
            bad.append((ref, e, v))
            continue
        tol = 1e-6 if abs(e) < 10 else 1.0
        d = abs(v - e)
        if d > tol:
            bad.append((ref, e, v))
        if abs(e) >= 10:
            worst = max(worst, d)
    elif v != e:
        bad.append((ref, e, v))
print(f"기대값 {len(exp)}개 대조 · 불일치 {len(bad)} · 최대 차이 {worst:.4f}원")
for b in bad[:15]:
    print("  ✗", b)
ok &= not bad

# (3) 금지 표현 — 내부 시스템·작성 경위·작업 메타
BANNED = [r"finModel", r"모델값", r"시스템 ?모델", r"사이트", r"화면", r"시트", r"대표 ?(지시|확정|결정|일정)", r"(?<![가-힣A-Za-z])형(?![가-힣])",
          r"(?<![가-힣])지시(?![가-힣])", r"adjust", r"costs\.py", r"정답지", r"레버", r"검증 ?보정", r"조사 ?(지적|권고)", r"판단값", r"\[(가정|추정|판단)\]",
          r"신뢰도", r"⚠", r"✓", r"ⓘ", r"\bIM\b", r"v\d\.\d", r"대조", r"GMV", r"5대 엔진", r"개시 계수", r"반영률", r"연 환산",
          r"salary_|mk\d|cpm_|ai_build", r"2026-09-1\d", r"법률의견", r"§", r"하이핀", r"하이가", r"온톨로지"]
hits = []
for ws in wv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("="):
                for pat in BANNED:
                    if re.search(pat, c.value):
                        hits.append((ws.title, c.coordinate, pat, c.value[:70]))
print(f"금지 표현 {len(hits)}건")
for h in hits[:25]:
    print("  ✗", h)
ok &= not hits

# (4) 인쇄 설정 — A4 가로 · 너비 맞춤
for ws in wf.worksheets:
    ps = ws.page_setup
    good = ps.orientation == "landscape" and str(ps.paperSize) == "9" and ws.sheet_properties.pageSetUpPr.fitToPage and ps.fitToWidth in (1, None)
    if not good:
        print("  ✗ 인쇄 설정", ws.title, ps.orientation, ps.paperSize, ps.fitToWidth)
    ok &= bool(good)
print("인쇄 설정 확인 완료")
print("결과:", "통과" if ok else "실패")
sys.exit(0 if ok else 1)
