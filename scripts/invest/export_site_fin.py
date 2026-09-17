# -*- coding: utf-8 -*-
"""예산양식(v3.x) → 사이트 재무회계 엔진 연동물 생성기.

1) src/data/finBudgetParams.js — 사이트 엔진(finBudget.js)이 읽는 입력(adjust.apply(_fin.json P) · costs.COST · 현금 레버 A/B/C · 투자 조건)
2) scripts/fin_site_golden.json — 사이트 엔진 회귀 대조용 정답값
   · annual: oracle.annual(P0) 연차별 계정(평탄화)
   · runs: oracle3.run(P0, LA/LB/LC) 저점·필요·요청·법인세·월별 유입/유출/누적·사용료 현금 인식 + 트랜치(T1·T2·T1 소진)
   · excel: 재계산 완료 엑셀의 연간손익(YR)·월별예산(MR 60개월)·현금(RA/RB/RC 84개월)·월별자금필요표(FR 39개월)·인력계획(HR)

사용: py scripts/invest/export_site_fin.py <재계산된 예산양식.xlsx>
엑셀 파이프라인(verify_budget.py 통과본)이 원천이다. 사이트 파라미터는 여기서만 만든다(손으로 고치지 않는다).
"""
import io
import json
import math
import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
sys.path.insert(0, HERE)
import adjust  # noqa: E402
import costs  # noqa: E402
import oracle as O  # noqa: E402
import oracle3 as O3  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

XLSX = sys.argv[1]
VERSION = os.path.basename(XLSX).rsplit("_v", 1)[-1].replace(".xlsx", "") if "_v" in XLSX else "v3.x"
if not VERSION.startswith("v"):
    VERSION = "v" + VERSION

P0 = O3.P0


def plain(v):
    """JSON으로 옮길 수 있게 튜플·중첩을 정리"""
    if isinstance(v, dict):
        return {str(k): plain(x) for k, x in v.items()}
    if isinstance(v, (list, tuple)):
        return [plain(x) for x in v]
    if isinstance(v, float) and (math.isinf(v) or math.isnan(v)):
        return None
    return v


def flat(d, pre=""):
    out = {}
    for k, v in d.items():
        key = f"{pre}{k}"
        if isinstance(v, dict):
            out.update(flat(v, key + "."))
        elif isinstance(v, (list, tuple)):
            for i, x in enumerate(v):
                if isinstance(x, (int, float)):
                    out[f"{key}.{i}"] = x
        elif isinstance(v, (int, float)) and not isinstance(v, bool):
            out[key] = v
    return out


# ── 1) 파라미터 ──
levers = {"A": O3.LA, "B": O3.LB, "C": O3.LC}
meta = {
    "version": VERSION,
    "source": "scripts/invest (adjust.py · costs.py · oracle.py · oracle3.py) — export_site_fin.py 생성물, 손으로 고치지 말 것",
    "calendarStart": "2027-01", "prepStart": "2026-10",
    "t1Rate": adjust.ADJ.get("t1Rate", 0.6), "imAmt": adjust.ADJ.get("imAmt", 0),
    "streams": [list(s) for s in adjust.STREAMS], "types": list(adjust.TYPES),
    "sections": [list(s[:4]) for s in costs.SECTIONS], "drivers": costs.DRIVERS,
    "burdenKeys": [list(b) for b in costs.BURDEN_KEYS], "channels": [list(c) for c in costs.CHANNELS],
    "aiModules": [list(m) for m in costs.AI_MODULES],
}
params_js = (
    "// 자동 생성 — scripts/invest/export_site_fin.py (예산양식 " + VERSION + ").  손으로 고치지 말 것: 예산양식 파이프라인에서 다시 생성한다.\n"
    "// 사이트 재무회계 엔진(finBudget.js)의 입력 = 투자금 산정 예산양식과 같은 가정값(finModel 기본값 + 대표 조정 adjust.py + 비용 근거 costs.py)\n"
    "const FB_P0 = " + json.dumps(plain(P0), ensure_ascii=False) + ";\n"
    "const FB_LEVERS = " + json.dumps(plain(levers), ensure_ascii=False) + ";\n"
    "const FB_META = " + json.dumps(plain(meta), ensure_ascii=False) + ";\n"
)
io.open(os.path.join(ROOT, "src", "data", "finBudgetParams.js"), "w", encoding="utf-8", newline="\n").write(params_js)

# ── 2) 정답값 ──
A = O.annual(P0)
annual = [flat(a) for a in A]


def tranche(req, rate, series):
    t1 = min(req, math.ceil(req * rate / 1e9) * 1e9)
    run_out = None
    for lab, idx, i_, o_, c_ in series:
        if t1 + c_ < 0:
            run_out = lab
            break
    return {"t1": t1, "t2": req - t1, "t1RunOut": run_out}


runs = {}
for nm, L in levers.items():
    r = O3.run(P0, L)
    runs[nm] = {
        "low": r["low"], "low_at": r["low_at"], "fixed_avg": r["fixed_avg"], "buffer": r["buffer"], "need": r["need"], "req": r["req"],
        "tax1": r["tax1"], "taxes": r["taxes"], "pre_exp": r["pre_exp"], "runway": r["runway"],
        "series": [list(s) for s in r["series"]], "insRecM": r["insRecM"],
        "tranche": tranche(r["req"], meta["t1Rate"], r["series"]),
    }

wb = load_workbook(XLSX, data_only=True)
mp = json.load(io.open(XLSX + ".map.json", encoding="utf-8"))


def rowvals(ws, row, c0, n):
    return [ws.cell(row=row, column=c0 + k).value for k in range(n)]


excel = {
    "YR": {k: rowvals(wb["연간손익"], r, 4, 5) for k, r in mp["YR"].items()},
    "MR": {k: rowvals(wb["월별예산"], r, 4, 60) for k, r in mp["MR"].items()},
    "RA": {k: rowvals(wb["현금·투자금"], r, 3, 84) for k, r in mp["RA"].items()},
    "RB": {k: rowvals(wb["현금·투자금"], r, 3, 84) for k, r in mp["RB"].items()},
    "RC": {k: rowvals(wb["현금·투자금"], r, 3, 84) for k, r in mp["RC"].items()},
    "FR": {k: rowvals(wb["월별자금필요표"], r, 4, 39) for k, r in mp["FR"].items()},
    "HR": {k: rowvals(wb["인력계획"], r, 4, 5) for k, r in mp["HR"].items()},
    "TR": {k: [wb["투자조건"].cell(row=r, column=3).value, wb["투자조건"].cell(row=r, column=4).value] for k, r in mp["TR"].items()},
}
# 엑셀 투자조건(기본 선택 = B)과 파이썬 트랜치 대조
tb = runs["B"]["tranche"]
assert abs((excel["TR"]["t1"][0] or 0) - tb["t1"]) < 1, ("T1 불일치", excel["TR"]["t1"], tb)
golden = {"version": VERSION, "annual": annual, "runs": runs, "excel": excel}
io.open(os.path.join(ROOT, "scripts", "fin_site_golden.json"), "w", encoding="utf-8", newline="\n").write(json.dumps(plain(golden), ensure_ascii=False))
print("params", len(params_js) // 1024, "KB · golden 연간", len(annual), "· 현금", {k: (round(v["req"] / 1e8), v["low_at"], v["tranche"]["t1RunOut"]) for k, v in runs.items()})
