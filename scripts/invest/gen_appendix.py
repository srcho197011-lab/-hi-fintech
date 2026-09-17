# -*- coding: utf-8 -*-
"""사업계획서 부록 — 「소요자금 산정 및 5개년 추정 재무계획(2027~2031)」 엑셀 생성기

투자금 산정 예산양식(엑셀 재계산 완료본)의 값을 읽어, 인쇄용(A4 가로) 부록 워크북을 만든다.
 - 원천 값(가정·계정 잎 행)은 값으로, 합계·이익·비율·잔액은 수식으로 넣는다.
 - 부록에 넣은 모든 수치 셀의 기대값을 <out>.expect.json에 남겨 check_appendix.py가 재계산 결과와 대조한다.
 - 문장 근거는 basis JSON(주요 가정 id별 문장 · 산정 근거 표)을 읽는다.

사용: py gen_appendix.py <원본 예산양식.xlsx> <map.json> <basis.json> <out.xlsx>
"""
import io
import json
import math
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break, RowBreak

SRC, MAP, BASIS, OUT = sys.argv[1:5]
wv = load_workbook(SRC, data_only=True)
mp = json.load(io.open(MAP, encoding="utf-8"))
BS = json.load(io.open(BASIS, encoding="utf-8"))
YR, AR, HR, FR, TR, CR = mp["YR"], mp["AR"], mp["HR"], mp["FR"], mp["TR"], mp["CR"]
Y, A, H, F, T, C = (wv[n] for n in ("연간손익", "가정", "인력계획", "월별자금필요표", "투자조건", "현금·투자금"))
MP = wv["매출계획"]

# ── 원천 정합 확인(행 번호표가 이 파일과 맞는지) ──
assert Y[f"B{YR['rev']}"].value == "매출액 합계", "연간손익 행 번호 불일치"
assert H[f"B{HR['payTotal']}"].value.startswith("인건비 합계"), "인력계획 행 번호 불일치"
assert F["D4"].value == "A 계획", "월별자금필요표 기준이 A 계획이 아님"
assert T["C5"].value == "B 보수", "투자조건 기준이 B 보수가 아님"
YC = "DEFGH"


def yv(k):
    return [Y[f"{c}{YR[k]}"].value or 0 for c in YC]


def av(k):
    return [A[f"{c}{AR[k]}"].value for c in YC]


def a1(k):
    return A[f"D{AR[k]}"].value


def hv(k):
    return [H[f"{c}{HR[k]}"].value or 0 for c in YC]


def cv(k):                                   # 현금·투자금 산정 결과 — A·B·C
    return [C[f"{c}{CR[k]}"].value for c in "CDE"]


# ══════════════════════════ 서식 ══════════════════════════
FN = "맑은 고딕"
NAVY, INK, SUBC, LINE, HDR, TOTF, KEYF = "1A2B4A", "111827", "4B5563", "C9D1DD", "E3E9F2", "F3F5F9", "FFF6D6"
MFMT = '#,##0,,;△#,##0,,;"-"'               # 원 값을 백만원으로 표시
NFMT = '#,##0;△#,##0;"-"'
PFMT = '0.0%;△0.0%;"-"'
WFMT = '#,##0;△#,##0;"-"'
D1 = '#,##0.0;△#,##0.0;"-"'
thin = Side(style="thin", color=LINE)
hair = Side(style="hair", color=LINE)
EXP = {}                                     # (시트, 셀) → 기대값


def font(size=9, bold=False, color=INK):
    return Font(name=FN, size=size, bold=bold, color=color)


def put(ws, ref, v, fmt=None, bold=False, size=9, color=INK, fill=None, h="left", wrap=False, exp=None, indent=0):
    c = ws[ref]
    c.value = v
    c.font = font(size, bold, color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=indent)
    if exp is not None:
        EXP[f"{ws.title}!{ref}"] = exp
    return c


def setup(ws, widths, title_rows=None, portrait=False):
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.sheet_view.showGridLines = False
    ps = ws.page_setup
    ps.orientation = "portrait" if portrait else "landscape"
    ps.paperSize = ws.PAPERSIZE_A4
    ps.fitToWidth = 1
    ps.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.35
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = ws.page_margins.footer = 0.25
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = "하이젠케어(주) 사업계획서 부록 — 소요자금 산정 및 5개년 추정 재무계획"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.font = FN
    ws.oddFooter.right.text = "&P / &N"
    ws.oddFooter.right.size = 8
    if title_rows:
        ws.print_title_rows = title_rows


def title(ws, r, text, sub=None, ncol=10):
    put(ws, f"A{r}", text, size=15, bold=True, color=NAVY)
    ws.row_dimensions[r].height = 26
    if sub:
        put(ws, f"A{r+1}", sub, size=9, color=SUBC)
    for i in range(1, ncol + 1):
        ws.cell(row=r + (2 if sub else 1), column=i).border = Border(top=Side(style="medium", color=NAVY))
    return r + (3 if sub else 2)


def section(ws, r, text, ncol):
    for i in range(1, ncol + 1):
        cc = ws.cell(row=r, column=i)
        cc.fill = PatternFill("solid", fgColor=NAVY)
    put(ws, f"A{r}", text, size=10, bold=True, color="FFFFFF", fill=NAVY)
    ws.row_dimensions[r].height = 19
    return r + 1


TEXT_HEADS = {"비고", "산정 근거", "산정 방식", "내용", "증원 기준(2028년부터)", "기준", "근거", "설명", "계정"}


def header(ws, r, labels, start=1, aligns=None):
    for i, lab in enumerate(labels):
        col = get_column_letter(start + i)
        h = (aligns[i] if aligns else ("left" if i < 2 or lab in TEXT_HEADS else "center"))
        put(ws, f"{col}{r}", lab, bold=True, color=NAVY, fill=HDR, h=h, wrap=True)
        ws[f"{col}{r}"].border = Border(bottom=thin, top=thin)
    ws.row_dimensions[r].height = 18 if max(len(str(x)) for x in labels) <= 9 else 28
    return r + 1


def hmerge(ws, r, c1, c2):
    """머리글 행의 c1~c2열을 합친다(채움·테두리는 범위 전체에)"""
    for i in range(c1, c2 + 1):
        cc = ws.cell(row=r, column=i)
        cc.fill = PatternFill("solid", fgColor=HDR)
        cc.border = Border(bottom=thin, top=thin)
    ws.cell(row=r, column=c1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)


def rule(ws, r, ncol, total=False):
    for i in range(1, ncol + 1):
        cc = ws.cell(row=r, column=i)
        cc.border = Border(bottom=thin if total else hair)
        if total:
            cc.fill = PatternFill("solid", fgColor=TOTF)


def note(ws, r, text, ncol, size=8.5, color=SUBC, height=None):
    put(ws, f"A{r}", text, size=size, color=color, wrap=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    if height:
        ws.row_dimensions[r].height = height
    return r + 1


def textw(s):
    """인쇄 폭 추정 — 한글 1, 공백 0.4, 그 밖의 문자 0.6"""
    return sum(1.0 if ord(ch) >= 0x1100 else (0.4 if ch == " " else 0.6) for ch in s)


def basis(key):
    return BS["assumptions"].get(key, "")


wb = Workbook()
wb.remove(wb.active)

# ══════════════════════════ 표지 ══════════════════════════
S0 = wb.create_sheet("표지")
setup(S0, [3, 26, 70, 26], portrait=False)
r = 3
put(S0, f"B{r}", "하이젠케어(주) 사업계획서", size=11, color=SUBC)
r += 1
put(S0, f"B{r}", "부록. 소요자금 산정 및 5개년 추정 재무계획", size=22, bold=True, color=NAVY)
S0.row_dimensions[r].height = 36
r += 1
put(S0, f"B{r}", "2027~2031년 · 작성 기준일 2026년 9월", size=11, color=SUBC)
for i in range(2, 5):
    S0.cell(row=r + 1, column=i).border = Border(top=Side(style="medium", color=NAVY))
r += 3
put(S0, f"B{r}", "목차", size=12, bold=True, color=NAVY)
r += 1
TOC = [("1. 요약", "투자 소요자금, 5개년 주요 지표, 자금 사용 계획, 분할 투자"),
       ("2. 주요 가정", "회원·제휴 기관·단가·비용 단가 등 추정에 사용한 가정과 산정 근거"),
       ("3. 매출 계획", "매출 항목별 산정 방식, 연도별 매출, 연말 규모 대비 실제 매출 비율, 2027년 월별 매출"),
       ("4. 추정 손익계산서", "2027~2031년 매출·원가·판매관리비·영업이익·당기순이익"),
       ("5. 인력 계획", "부문별 인원과 인건비, 인건비 산정 기준"),
       ("6. 월별 자금 계획", "2026년 10월~2029년 12월 월별 매출·입금·지출·누적 현금(기본 계획)"),
       ("7. 투자금 산정", "산정 원칙, 계획별(기본·보수·지연) 소요자금, 현금 흐름 가정, 법인세"),
       ("8. 단가 산정 근거", "비용·매출 단가의 조사 근거와 출처")]
for a_, b_ in TOC:
    put(S0, f"B{r}", a_, size=10, bold=True)
    put(S0, f"C{r}", b_, size=10, color=SUBC)
    for i in range(2, 5):
        S0.cell(row=r, column=i).border = Border(bottom=hair)
    S0.row_dimensions[r].height = 18
    r += 1
r += 1
put(S0, f"B{r}", "작성 기준", size=12, bold=True, color=NAVY)
r += 1
BASE = [
    ("사업 일정", "2026년 10월 사업 준비 착수, 2027년 1월 서비스 개시"),
    ("추정 기간", "손익 2027~2031년(5개년, 달력 연도) · 월별 자금 2026년 10월~2029년 12월"),
    ("금액 단위", "백만원(부가가치세 별도) · 단가는 원 · 음수는 △로 표기"),
    ("매출 계산", "연중 회원과 제휴 기관이 늘어나는 흐름에 맞춰 월별로 누적 계산"),
    ("투자금 산정", "월별 누적 현금의 최저점과 6개월분 고정비를 더한 금액을 10억원 단위로 올림"),
]
for a_, b_ in BASE:
    put(S0, f"B{r}", a_, size=10, bold=True)
    put(S0, f"C{r}", b_, size=10, wrap=True)
    S0.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
    S0.row_dimensions[r].height = 20
    r += 1

# ══════════════════════════ 4. 추정 손익계산서(다른 시트가 참조하므로 먼저 만든다) ══════════════════════════
S4 = wb.create_sheet("4.추정손익")
W4 = [9, 32, 38, 12, 12, 12, 12, 12, 13]
setup(S4, W4, title_rows="4:4")
r = title(S4, 1, "4. 추정 손익계산서(2027~2031)", "단위: 백만원 · 운영 지표는 명·건·곳", 9)
HDR4 = ["구분", "계정", "산정 방식", "2027년", "2028년", "2029년", "2030년", "2031년", "5년 합계"]
r = header(S4, 4, HDR4)
R4 = {}


def row4(key, grp, label, calc, vals=None, formula=None, fmt=MFMT, total=False, sum5=True, expect=None, bold=None):
    """vals: 5개 값 · formula: i -> 수식 문자열(i=0..4)"""
    global r
    put(S4, f"A{r}", grp, bold=True, color=SUBC)
    put(S4, f"B{r}", label, bold=bool(total if bold is None else bold), indent=0 if total else 1)
    put(S4, f"C{r}", calc, size=8.5, color=SUBC, wrap=True)
    for i in range(5):
        col = YC[i]
        e = (expect[i] if expect is not None else (vals[i] if vals is not None else None))
        v = formula(i) if formula else vals[i]
        put(S4, f"{col}{r}", v, fmt=fmt, bold=total, h="right", exp=e)
    if sum5:
        e5 = sum(expect) if expect is not None else (sum(vals) if vals is not None else None)
        put(S4, f"I{r}", f"=SUM(D{r}:H{r})", fmt=fmt, bold=True, h="right", exp=e5)
    rule(S4, r, 9, total)
    R4[key] = r
    r += 1


def sumf(*keys):
    return lambda i: "=" + "+".join(f"{YC[i]}{R4[k]}" for k in keys)


def rngf(k1, k2):
    return lambda i: f"=SUM({YC[i]}{R4[k1]}:{YC[i]}{R4[k2]})"


r = section(S4, r, "주요 운영 지표", 9)
row4("me", "회원", "연말 회원", "연말 누적(이탈 반영)", yv("me"), fmt=NFMT, sum5=False)
row4("new", "", "순증 회원", "연말 − 전년 말", yv("new"), fmt=NFMT)
row4("gnew", "", "신규 가입 회원", "순증 + 전년 말 회원 × 이탈률 18%", yv("gross_new"), fmt=NFMT)
row4("chk", "이용", "건강검진 예약", "월말 회원 × 예약 비율 × 계절 지수의 월별 합계", yv("activeAct"), fmt=NFMT)
row4("paid", "", "AI 플랫폼 구독 기관(연말)", "검진센터 + 병원 + 약국 · 검진센터는 2028년부터 과금", [a + b + c for a, b, c in zip(yv("paid_c"), yv("paid_h"), yv("paid_p"))], fmt=NFMT, sum5=False)
row4("ins", "", "헬스메이트센터 DB 공급", "월말 누적 동의 회원 × 60% ÷ 12의 월별 합계(2027년은 공급 계획 12만건)", yv("insCases"), fmt=NFMT)
row4("buyers", "", "건강커머스 구매 회원", "연말 회원 × 38%", yv("buyers"), fmt=NFMT, sum5=False)
row4("head", "인력", "인원(연말, 자문위원 제외)", "5. 인력 계획", yv("headTotal"), fmt=NFMT, sum5=False)

r = section(S4, r, "매출액", 9)
CATS = [("supp", "영양제·보충제"), ("diet", "건강식단·식품"), ("device", "홈케어 의료기기"), ("sports", "스포츠용품·활동")]
for k, lab in CATS:
    row4("rev_" + k, "건강커머스" if k == "supp" else "", f"상품 판매 — {lab}", "월말 구매 회원 × 1인 연 구매액 ÷ 12 × 70% × 적용률의 월별 합계" if k == "supp" else "", yv("rev_" + k))
row4("revP", "", "건강커머스 소계", "", formula=rngf("rev_supp", "rev_sports"), expect=yv("revP"), total=True)
row4("revChk_o", "건강검진", "건강검진 연계 — 자사 운영", "자사 건수 × 건당 3만원", yv("revChk_o"))
row4("revChk_t", "", "건강검진 연계 — 제휴사", "제휴사 건수 × 건당 1만원", yv("revChk_t"))
row4("revChk", "", "건강검진 연계 소계", "", formula=sumf("revChk_o", "revChk_t"), expect=yv("revChk"), total=True)
row4("revResv", "예약", "예약 서비스 수수료(골프·시설)", "월별 예약 건수 합계 × 건당 1만원(계절 지수 미적용)", yv("revResv"))
row4("revCare", "돌봄", "재가·돌봄 파트너 이용료", "과금 센터·월 × 월 30만원", yv("revCare"))
row4("sub_c", "구독", "AI 플랫폼 구독 — 검진센터", "월별 과금 기관 × 월 구독료의 월별 합계", yv("sub_c"))
row4("sub_h", "", "AI 플랫폼 구독 — 병원", "", yv("sub_h"))
row4("sub_p", "", "AI 플랫폼 구독 — 약국", "", yv("sub_p"))
row4("revSub", "", "AI 플랫폼 구독 소계", "", formula=rngf("sub_c", "sub_p"), expect=yv("revSub"), total=True)
row4("revIns_d", "사용료", "헬스메이트센터 사용료 — 우대 단가분", "우대 한도 내 건수 × 5만원", yv("revIns_d"))
row4("revIns_f", "", "헬스메이트센터 사용료 — 시가분", "한도 초과 건수 × 10만원", yv("revIns_f"))
row4("revIns", "", "헬스메이트센터 사용료 소계", "", formula=sumf("revIns_d", "revIns_f"), expect=yv("revIns"), total=True)
row4("rev", "", "매출액 합계", "", formula=sumf("revP", "revChk", "revResv", "revCare", "revSub", "revIns"), expect=yv("rev"), total=True)

r = section(S4, r, "매출원가", 9)
for k, lab in CATS:
    row4("cogs_" + k, "상품" if k == "supp" else "", f"상품 원가 — {lab}", "품목 매출 × 원가율(35~60%)" if k == "supp" else "", yv("cogs_" + k))
row4("cogsP", "", "상품 원가 소계", "", formula=rngf("cogs_supp", "cogs_sports"), expect=yv("cogsP"), total=True)
row4("chkCogs_o", "검진", "건강검진 연계 서비스 원가 — 자사 운영", "자사 건수 × 건당 2만원", yv("chkCogs_o"))
row4("chkCogs_t", "", "건강검진 연계 서비스 원가 — 제휴사", "제휴사 건수 × 건당 2만원", yv("chkCogs_t"))
row4("chkCogs", "", "건강검진 연계 서비스 원가 소계", "", formula=sumf("chkCogs_o", "chkCogs_t"), expect=yv("chkCogs"), total=True)
row4("careCost", "기타", "재가·돌봄 연계 운영 원가", "돌봄 매출 × 5%", yv("careCost"))
row4("subCost", "", "구독 운영 원가(클라우드·연동)", "구독 매출 × 12%", yv("subCost"))
row4("payFee", "", "결제 대행 수수료", "건강커머스 매출 × 2.2%", yv("payFee"))
row4("cogs", "", "매출원가 합계", "", formula=sumf("cogsP", "chkCogs", "careCost", "subCost", "payFee"), expect=yv("cogs"), total=True)
row4("gross", "", "매출총이익", "매출액 − 매출원가", formula=lambda i: f"={YC[i]}{R4['rev']}-{YC[i]}{R4['cogs']}", expect=yv("gross"), total=True)
row4("gm", "", "매출총이익률", "", formula=lambda i: f"=IF({YC[i]}{R4['rev']}=0,0,{YC[i]}{R4['gross']}/{YC[i]}{R4['rev']})", expect=yv("gm"), fmt=PFMT, sum5=False, bold=False)

r = section(S4, r, "판매관리비", 9)
row4("mk1", "마케팅", "메디에이지 검진 시기 안내 발송", "분기 150만명 × 3회 × 4분기 × 31원", yv("mk1"))
row4("media", "", "온라인 타겟 광고(8개 채널)", "월 노출 × 12 × 가중 CPM ÷ 1,000", yv("media"))
row4("creative", "", "광고 소재 제작", "영상 4편 × 300만원 + 카드뉴스 12세트 × 15만원", yv("creative"))
row4("kit", "", "검진센터 QR 안내물", "신규 배치 센터 × 6만원", yv("kit"))
row4("sticker", "", "결과지 봉투 QR 스티커", "검진 예약 × 매당 25원", yv("sticker"))
row4("mktSum", "", "마케팅 소계", "", formula=rngf("mk1", "sticker"), expect=yv("mktSum"), total=True)
row4("mediRep", "제휴", "메디에이지 제휴 리포트 구매", "2027년 가입 회원 33만명 × 건당 3,000원", yv("mediRep"))
row4("reward", "고객", "회원 포인트 적립", "상품 마진 × 60%", yv("reward"))
row4("donation", "", "기부금(치료비 나눔)", "상품 마진 × 15%", yv("donation"))
row4("pay", "인력", "인건비", "5. 인력 계획", yv("pay"))
row4("itMaint", "IT 운영", "AI 시스템 유지보수", "누적 구축비(보안 인증·사무 장비 제외) × 연 15%(당해 투자분은 절반)", yv("itMaint"))
row4("itData", "", "데이터 유지관리", "월 400만원 × 12 × √(연말 회원 ÷ 2027년 연말 회원)", yv("itData"))
row4("itSec", "", "보안관제", "월 120만원 × 12", yv("itSec"))
row4("cloudBase", "", "클라우드 기본 환경", "월 380만원 × 12", yv("cloudBase"))
row4("cloudVar", "", "클라우드 증설", "30만명 초과 평균 회원 × 월 12원 × 12", yv("cloudVar"))
row4("llm", "", "AI 상담 이용료(LLM API)", "평균 회원 × 연 12회 × 토큰 단가", yv("llm"))
row4("bc", "", "블록체인 기록 비용", "기록 건수 × 건당 40원", yv("bc"))
row4("itOpex", "", "IT 운영비 소계", "", formula=rngf("itMaint", "bc"), expect=yv("itOpex"), total=True)
row4("salesCost", "운영", "영업비", "매출 × 2% × 30%", yv("salesCost"))
row4("adminCost", "", "관리비", "매출 × 4% × 30%", yv("adminCost"))
row4("sga", "", "판매관리비 합계", "", formula=sumf("mktSum", "mediRep", "reward", "donation", "pay", "itOpex", "salesCost", "adminCost"), expect=yv("sga"), total=True)

S4.row_breaks.append(Break(id=r - 1))
r = section(S4, r, "손익", 9)
# 이월결손금 = 2026년 10~12월 지출 − 시스템 구축비(자산) − 임차 보증금
out26 = sum(F[f"{get_column_letter(4 + k)}{FR['oTot']}"].value or 0 for k in range(3))
build26 = sum(F[f"{get_column_letter(4 + k)}{FR['oBuild']}"].value or 0 for k in range(3))
nol27 = yv("nolOpen")[0]
dep26 = out26 - build26 - nol27
assert 0 < dep26 < 2e8, f"준비기간 보증금 추정 이상: {dep26}"
eok = lambda x: f"{x / 1e8:.1f}억원"
row4("ebitda", "", "상각 전 영업이익(EBITDA)", "매출총이익 − 판매관리비", formula=lambda i: f"={YC[i]}{R4['gross']}-{YC[i]}{R4['sga']}", expect=yv("ebitModel"), total=True)
row4("depr", "", "감가상각비", "투자액 ÷ 5년 정액(취득 연도 반년)", yv("depr"))
row4("ebit", "", "영업이익", "상각 전 영업이익 − 감가상각비", formula=lambda i: f"={YC[i]}{R4['ebitda']}-{YC[i]}{R4['depr']}", expect=yv("ebit"), total=True)
row4("opm", "", "영업이익률", "", formula=lambda i: f"=IF({YC[i]}{R4['rev']}=0,0,{YC[i]}{R4['ebit']}/{YC[i]}{R4['rev']})", expect=yv("opm"), fmt=PFMT, sum5=False, bold=False)
row4("pbt", "", "법인세비용차감전순이익", "영업이익(차입금 없음)", formula=lambda i: f"={YC[i]}{R4['ebit']}", expect=yv("pbt"), total=True)
row4("nol", "", "이월결손금 공제", f"2026년 준비기간 비용(지출 {eok(out26)} 중 시스템 구축비 {eok(build26)}·보증금 {eok(dep26)} 제외)", yv("nolOpen"))
row4("tax", "", "법인세", "(세전이익 − 이월결손금) × 22%", formula=lambda i: f"=MAX(0,{YC[i]}{R4['pbt']}-{YC[i]}{R4['nol']})*0.22", expect=yv("tax"))
row4("net", "", "당기순이익", "세전이익 − 법인세", formula=lambda i: f"={YC[i]}{R4['pbt']}-{YC[i]}{R4['tax']}", expect=yv("net"), total=True)

r = section(S4, r, "참고 — 투자", 9)
row4("capex", "", "시스템 구축·고도화·장비 투자", "2. 주요 가정 ⑩", yv("capex"))
r = note(S4, r + 1, "※ 손실 연도는 법인세 없음. 2026년 준비기간 비용을 이월결손금으로 공제, 현금 납부는 다음 해 3월", 9)

# ══════════════════════════ 1. 요약 ══════════════════════════
S1 = wb.create_sheet("1.요약", 1)
W1 = [46, 13, 13, 13, 13, 13, 14, 40]
setup(S1, W1)
r = title(S1, 1, "1. 요약", "단위: 백만원 · 투자 요청액은 10억원 단위", 8)
r = section(S1, r, "① 투자 소요자금 — 계획별", 8)
r = header(S1, r, ["구분", "A 기본 계획", "B 보수 계획", "C 지연 계획", "비고", "", "", ""],
           aligns=["left", "center", "center", "center", "left", "left", "left", "left"])
hmerge(S1, r - 1, 5, 8)
SUMROW = {}
low, lowcal, fixed, buf, need, req = cv("s_low"), cv("s_lowCal"), cv("s_fixed"), cv("s_buffer"), cv("s_need"), cv("s_req")
lowcal = [f"{v[:4]}년 {int(v[5:7])}월" for v in lowcal]
lines = [("low", "누적 현금 최저점(투자 전)", low, MFMT, "투자 없이 운영할 때 가장 크게 모자라는 금액"),
         ("lowcal", "최저점 발생 월", lowcal, None, ""),
         ("buf", "안전 버퍼", buf, MFMT, "2027년 월평균 고정비 6개월분"),
         ("need", "필요 총자금", None, MFMT, "최저점 부족액 + 안전 버퍼"),
         ("req", "투자 요청액", None, MFMT, "필요 총자금을 10억원 단위로 올림")]
for key, lab, vals, fmt, rem in lines:
    put(S1, f"A{r}", lab, bold=key in ("need", "req"))
    for i, col in enumerate("BCD"):
        if key == "need":
            put(S1, f"{col}{r}", f"=-{col}{SUMROW['low']}+{col}{SUMROW['buf']}", fmt=fmt, bold=True, h="right", exp=need[i])
        elif key == "req":
            put(S1, f"{col}{r}", f"=ROUNDUP({col}{SUMROW['need']}/1000000000,0)*1000000000", fmt=fmt, bold=True, h="right", exp=req[i], fill=KEYF)
        else:
            put(S1, f"{col}{r}", vals[i], fmt=fmt, h="right" if fmt else "center", exp=(vals[i] if fmt else None))
    put(S1, f"E{r}", rem, size=8.5, color=SUBC, indent=1)
    rule(S1, r, 8, key in ("need", "req"))
    S1.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    SUMROW[key] = r
    r += 1
for t_ in ["A 기본 계획: 2027년 1월 서비스 개시",
           "B 보수 계획: 개시 2개월 지연, 헬스메이트센터 사용료 입금 2개월 지연, 시스템 고도화 투자를 분기 초에 미리 집행",
           "C 지연 계획: B 조건에 더해 개시 후 12개월간 헬스메이트센터 DB 공급 없음(7. 투자금 산정 참조)"]:
    r = note(S1, r, t_, 8)
r += 1

r = section(S1, r, "② 5개년 주요 지표", 8)
r = header(S1, r, ["구분", "2027년", "2028년", "2029년", "2030년", "2031년", "5년 합계", "비고"], aligns=["left"] + ["center"] * 6 + ["left"])
KP = [("연말 회원(명)", "me", NFMT, False, ""), ("매출액", "rev", MFMT, True, ""), ("매출총이익", "gross", MFMT, True, ""),
      ("매출총이익률", "gm", PFMT, False, ""), ("영업이익", "ebit", MFMT, True, ""), ("영업이익률", "opm", PFMT, False, ""),
      ("당기순이익", "net", MFMT, True, ""), ("인원(명)", "head", NFMT, False, "자문위원 제외")]
for lab, k, fmt, s5, rem in KP:
    put(S1, f"A{r}", lab, bold=k in ("rev", "ebit"))
    src_row = R4[k]
    for i, col in enumerate("BCDEF"):
        exp_v = S4[f"{YC[i]}{src_row}"].value
        exp_v = EXP.get(f"4.추정손익!{YC[i]}{src_row}", exp_v)
        put(S1, f"{col}{r}", f"='4.추정손익'!{YC[i]}{src_row}", fmt=fmt, h="right", exp=exp_v)
    if s5:
        put(S1, f"G{r}", f"=SUM(B{r}:F{r})", fmt=fmt, bold=True, h="right", exp=EXP.get(f"4.추정손익!I{src_row}"))
    put(S1, f"H{r}", rem, size=8.5, color=SUBC, indent=1)
    rule(S1, r, 8)
    r += 1
r += 1

S1.row_breaks.append(Break(id=r - 1))
r = section(S1, r, f"③ 자금 사용 계획 — B 보수 계획 기준(투자 요청액 {req[1] / 1e8:,.0f}억원)", 8)
r = header(S1, r, ["사용처", "금액", "비중", "산정 방식", "", "", "", ""], aligns=["left", "right", "right", "left", "left", "left", "left", "left"])
hmerge(S1, r - 1, 4, 8)
USE_NOTE = {}
LABEL = {"매출원가": "매출원가(건강검진 연계 서비스·상품·결제 수수료)", "포인트 적립 · 기부금": "회원 포인트 적립·기부금",
         "마케팅(5대 엔진 광고 · 출시 전 사전 광고)": "마케팅(회원 확보 광고·개시 전 사전 광고)",
         "CAPEX(1차 시스템 설치 · 고도화 · 보안 · 장비)": "시스템 구축·고도화·보안·장비",
         "메디에이지 리포트 구매(제휴 DB)": "메디에이지 제휴 리포트 구매", "인건비(준비기간 포함)": "인건비(준비기간 포함)",
         "AI·데이터·클라우드 · 영업 · 관리(준비기간 IT·기타 포함 · 사용료 미공급 절감 차감)": "IT 운영·영업·관리비",
         "임차 · 보증금 · 채용 · 일회성": "임차료·보증금·채용·설립 비용"}
use_rows = []
for rr in range(TR["catStart"], TR["catEnd"] + 1):
    lab, amt = T[f"B{rr}"].value, T[f"E{rr}"].value or 0
    if lab.startswith("이자"):
        assert abs(amt) < 1, "이자·세금 항목이 0이 아님"
        continue
    put(S1, f"A{r}", LABEL.get(lab, lab), indent=1)
    put(S1, f"B{r}", amt, fmt=MFMT, h="right", exp=amt)
    use_rows.append(r)
    rule(S1, r, 8)
    r += 1
first_use, last_use = use_rows[0], use_rows[-1]
put(S1, f"A{r}", "소계 — 최저점까지 순부족액", bold=True)
put(S1, f"B{r}", f"=SUM(B{first_use}:B{last_use})", fmt=MFMT, bold=True, h="right", exp=T[f"E{TR['outSum'] + 2}"].value)
USE_NOTE[r] = "최저점 달까지의 사용처별 누적 지출 비중으로 최저점 부족액을 나눈 금액"
rule(S1, r, 8, True)
sub_r = r
r += 1
put(S1, f"A{r}", "안전 버퍼", indent=1)
put(S1, f"B{r}", T[f"E{TR['outSum'] + 3}"].value, fmt=MFMT, h="right", exp=T[f"E{TR['outSum'] + 3}"].value)
USE_NOTE[r] = "2027년 월평균 고정비 6개월분"
rule(S1, r, 8)
buf_r = r
r += 1
put(S1, f"A{r}", "10억원 단위 올림", indent=1)
put(S1, f"B{r}", T[f"E{TR['outSum'] + 5}"].value, fmt=MFMT, h="right", exp=T[f"E{TR['outSum'] + 5}"].value)
rule(S1, r, 8)
rnd_r = r
r += 1
put(S1, f"A{r}", "합계(투자 요청액)", bold=True)
put(S1, f"B{r}", f"=B{sub_r}+B{buf_r}+B{rnd_r}", fmt=MFMT, bold=True, h="right", fill=KEYF, exp=T[f"E{TR['useTotal']}"].value)
rule(S1, r, 8, True)
tot_r = r
for rr in use_rows + [sub_r, buf_r, rnd_r, tot_r]:
    put(S1, f"C{rr}", f"=IF($B${tot_r}=0,0,B{rr}/$B${tot_r})", fmt=PFMT, h="right")
    put(S1, f"D{rr}", USE_NOTE.get(rr, ""), size=8.5, color=SUBC, indent=1)
    S1.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=8)
r += 2

r = section(S1, r, "④ 분할 투자 — B 보수 계획 기준", 8)
r = header(S1, r, ["구분", "금액", "내용", "", "", "", "", ""], aligns=["left", "right", "left", "left", "left", "left", "left", "left"])
hmerge(S1, r - 1, 3, 8)
t1, t2 = T[f"C{TR['t1']}"].value, T[f"C{TR['t1'] + 1}"].value
t1_until = T[f"D{TR['t1run']}"].value
ty, tm = int(t1_until[:4]), int(t1_until[5:7])
dy, dm = (ty, tm - 1) if tm > 1 else (ty - 1, 12)
TR_ROWS = [("1차 투자(계약 시)", t1, f"요청액의 60%를 10억원 단위로 올린 금액 · 1차 투자금만으로는 {ty}년 {tm}월에 자금 소진"),
           ("2차 투자(성과 확인 후)", t2, f"아래 성과 목표 확인 후 {dy}년 {dm}월까지 납입")]
for lab_, amt_, txt_ in TR_ROWS:
    put(S1, f"A{r}", lab_, indent=1)
    put(S1, f"B{r}", amt_, fmt=MFMT, h="right", exp=amt_)
    put(S1, f"C{r}", txt_, size=8.5, color=SUBC, indent=1)
    S1.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    rule(S1, r, 8)
    r += 1
put(S1, f"A{r}", "합계", bold=True)
put(S1, f"B{r}", f"=B{r-2}+B{r-1}", fmt=MFMT, bold=True, h="right", exp=t1 + t2)
rule(S1, r, 8, True)
r += 1
MILE = ["헬스메이트센터 DB 공급 시범 운영 성과 확인(응답률 80%, 상담 전환 40%, 청약 전환 8%)",
        "보험업·신용정보·개인정보 관련 법률 검토 의견서 확보",
        "현대해상 고객의 제3자 제공 동의 수집 목표 달성",
        "AI 플랫폼 유료 구독 기관 목표 달성"]
put(S1, f"A{r}", "2차 투자 성과 목표", bold=True, color=NAVY)
r += 1
for i_, m_ in enumerate(MILE):
    r = note(S1, r, f"{i_ + 1}) {m_}", 8, size=9, color=INK)

# ══════════════════════════ 2. 주요 가정 ══════════════════════════
S2 = wb.create_sheet("2.주요가정", 2)
W2 = [12, 34, 10, 11, 11, 11, 11, 11, 62]
setup(S2, W2, title_rows="4:4")
r = title(S2, 1, "2. 주요 가정", "2027~2031년 열을 합친 한 칸 값은 전 기간 같은 값 적용 · 2027년 열에만 값이 있는 항목은 2027년 기준값(각 행 산정 근거 참조)", 9)
r = header(S2, 4, ["구분", "항목", "단위", "2027년", "2028년", "2029년", "2030년", "2031년", "산정 근거"])
r = 5


def arow(grp, label, unit, vals, fmt, bid=None, once=False):
    """vals: 5개 리스트(연도별) 또는 스칼라(전 기간 동일 — D:H 병합) · once=True면 2027년에만"""
    global r
    put(S2, f"A{r}", grp, bold=True, color=SUBC)
    put(S2, f"B{r}", label, indent=1)
    put(S2, f"C{r}", unit, size=8.5, color=SUBC, h="center")
    if isinstance(vals, list):
        for i in range(5):
            put(S2, f"{YC[i]}{r}", vals[i], fmt=fmt, h="right", exp=vals[i])
    elif once:
        put(S2, f"D{r}", vals, fmt=fmt, h="right", exp=vals)
    else:
        put(S2, f"D{r}", vals, fmt=fmt, h="center", exp=vals)
        S2.merge_cells(start_row=r, start_column=4, end_row=r, end_column=8)
    arow_basis(r, bid)
    rule(S2, r, 9)
    AROW[label] = r
    r += 1


def arow_basis(r_, bid):
    txt = basis(bid) if bid else ""
    put(S2, f"I{r_}", txt, size=8.5, color=INK, wrap=True)
    S2.row_dimensions[r_].height = 13 * max(1, math.ceil(textw(txt) / 46)) + 2    # 산정 근거 열 한 줄 약 46(한글 1자 기준)


def arow_ratio(grp, label, num_lab, den, bid, expect):
    """den: i -> 분모 수식 조각 · 결과는 수식(비율)"""
    global r
    put(S2, f"A{r}", grp, bold=True, color=SUBC)
    put(S2, f"B{r}", label, indent=1)
    put(S2, f"C{r}", "%", size=8.5, color=SUBC, h="center")
    for i in range(5):
        col = YC[i]
        put(S2, f"{col}{r}", f"=IF({den(i)}=0,0,{col}{AROW[num_lab]}/({den(i)}))", fmt=PFMT, h="right", exp=expect[i])
    arow_basis(r, bid)
    rule(S2, r, 9)
    r += 1


AROW = {}
r = section(S2, r, "① 회원·이용", 9)
arow("회원", "연말 회원", "명", av("membersEnd"), NFMT, "members")
arow("", "건강검진 예약(연간 규모)", "건", av("activeAbs"), NFMT, "active")
arow_ratio("", "회원당 연간 검진 예약 비율", "건강검진 예약(연간 규모)", lambda i: f"{YC[i]}{AROW['연말 회원']}", "activeRatio",
           [a / m for a, m in zip(av("activeAbs"), av("membersEnd"))])
arow("", "보험 안내 동의 회원(누적)", "명", av("mktConsent"), NFMT, "consent")
arow("", "연간 회원 이탈률", "%", a1("churn"), PFMT, "churn")
r = section(S2, r, "② 제휴 기관·AI 플랫폼 구독", 9)
arow("기관", "제휴 검진센터(연말)", "곳", av("centers"), NFMT, "centers")
arow("", "제휴 병원(연말)", "곳", av("hospitals"), NFMT, "hospitals")
arow("", "제휴 약국(연말)", "곳", av("pharmacies"), NFMT, "pharmacies")
arow("구독료", "검진센터 월 구독료", "원/월", yv("fee_c"), WFMT, "fee_c")
arow("", "병원 월 구독료", "원/월", yv("fee_h"), WFMT, "fee_h")
arow("", "약국 월 구독료", "원/월", yv("fee_p"), WFMT, "fee_p")
arow("", "구독 운영 원가율", "%", a1("subCost"), PFMT, "subCost")
r = section(S2, r, "③ 건강커머스", 9)
arow("구매", "구매 회원 비율(연말 회원 대비)", "%", a1("buyerRate"), PFMT, "buyerRate")
arow("", "당사 채널 구매 비중", "%", a1("capture"), PFMT, "capture")
for k, lab in CATS:
    arow("", f"1인 연 구매액 — {lab}", "원/년", a1("arpu_" + k), WFMT, "arpu" if k == "supp" else None)
for k, lab in CATS:
    arow("원가" if k == "supp" else "", f"상품 원가율 — {lab}", "%", a1("cost_" + k), PFMT, "cogs" if k == "supp" else None)
arow("", "결제 대행 수수료율", "%", a1("payRate"), PFMT, "payRate")
arow("고객", "회원 포인트 적립률(상품 마진 대비)", "%", a1("rewardRate"), PFMT, "reward")
arow("", "기부금 비율(상품 마진 대비)", "%", a1("donationRate"), PFMT, "donation")
arow("조정", "2027년 건강커머스 매출 적용률(30% 감액)", "%", a1("prodAdjY1"), PFMT, "prodAdj", once=True)
r = section(S2, r, "④ 건강검진 연계·예약 서비스·재가·돌봄", 9)
arow("검진", "건강검진 연계(자사 운영) — 건당 매출", "원/건", a1("chkOwnFee"), WFMT, "chkOwn")
arow("", "건강검진 연계(자사 운영) — 건당 원가", "원/건", a1("chkOwnCost"), WFMT)
arow("", "건강검진 연계(제휴사) — 건당 매출", "원/건", a1("chkPtnFee"), WFMT, "chkPtn")
arow("", "건강검진 연계(제휴사) — 건당 원가", "원/건", a1("chkPtnCost"), WFMT)
arow("", "자사 운영 비중", "%", yv("chkShare"), PFMT, "chkShare")
arow("예약", "검진 예약 1건당 예약 서비스 건수", "건", av("resvPer"), D1, "resvPer")
arow("", "예약 서비스 건당 수수료", "원/건", a1("resvFee"), WFMT, "resvFee")
arow("돌봄", "재가·돌봄 수요 발생률(연말 회원 대비)", "%", a1("careRate"), PFMT, "careRate")
arow("", "파트너 센터당 월 연결 건수", "건", a1("carePerCenter"), NFMT, "careCenter")
arow("", "파트너 센터 월 이용료", "원/월", a1("careFee"), WFMT, "careFee")
arow("", "재가·돌봄 연계 운영 원가율", "%", a1("careCost"), PFMT, "careCost")
r = section(S2, r, "⑤ 헬스메이트센터 사용료", 9)
arow("공급", "동의 DB 연간 공급률(누적 동의 대비)", "%", a1("insConv"), PFMT, "insConv")
arow("", "DB 공급 건수(연간)", "건", yv("insCases"), NFMT, "ins2027")
_conv_row = AROW["동의 DB 연간 공급률(누적 동의 대비)"]
arow_ratio("", "실제 공급 비율(연말 동의 × 60% 대비)", "DB 공급 건수(연간)",
           lambda i: f"{YC[i]}{AROW['보험 안내 동의 회원(누적)']}*$D${_conv_row}", "insShare",
           [n / (c * a1("insConv")) for n, c in zip(yv("insCases"), av("mktConsent"))])
arow("단가", "DB 건당 시가", "원/건", a1("hmMarket"), WFMT, "hmMarket")
arow("", "전략적 투자자 우대 단가", "원/건", yv("hmPrice"), WFMT, "hmPrice")
arow("", "우대 단가 적용 한도", "건", yv("hmCap"), NFMT, "hmCap")
r = section(S2, r, "⑥ 마케팅", 9)
arow("안내", "메디에이지 검진 시기 안내 대상(분기)", "명", av("mk1Target"), NFMT, "mk1")
arow("", "1인당 분기 발송 횟수", "회", a1("mk1Times"), NFMT)
arow("", "안내 발송 단가(LMS)", "원/건", a1("mk1Unit"), WFMT)
arow("광고", "온라인 타겟 광고 월 노출", "회/월", yv("mk2Impr"), NFMT, "mk2")
arow("", "채널 가중 평균 CPM", "원/천회", a1("mk2Cpm"), WFMT)
arow("", "영상 광고 소재 제작", "편/년", a1("videoN"), NFMT, "creative")
arow("", "영상 소재 1편 제작비", "원/편", a1("videoUnit"), WFMT)
arow("", "카드뉴스·배너 제작", "세트/년", a1("cardN"), NFMT)
arow("", "카드뉴스·배너 1세트 제작비", "원/세트", a1("cardUnit"), WFMT)
arow("현장", "QR 안내물 배치 검진센터(누적)", "곳", av("qrCenters"), NFMT, "qr")
arow("", "센터당 QR 안내물", "원/곳", a1("qrKit"), WFMT)
arow("", "결과지 봉투 QR 스티커", "원/매", a1("qrSticker"), WFMT)
r = section(S2, r, "⑦ 메디에이지 제휴 DB", 9)
arow("제휴", "제휴 DB 확보 규모", "건", a1("mediDb"), NFMT, "medi", once=True)
arow("", "확보 기간(2027년 1월부터)", "개월", a1("mediMonths"), NFMT, once=True)
arow("", "리포트 구매 건수(가입 회원)", "건", yv("mediN"), NFMT)
arow("", "리포트 구매 단가", "원/건", a1("mediFee"), WFMT)
r = section(S2, r, "⑧ IT 운영비", 9)
arow("시스템", "AI 시스템 유지보수(누적 구축비 대비, 연)", "%", a1("maintRate"), PFMT, "itMaint")
arow("", "데이터 유지관리", "원/월", a1("dataMonth"), WFMT, "itData")
arow("", "보안관제", "원/월", a1("secMonth"), WFMT, "itSec")
arow("클라우드", "클라우드 기본 환경", "원/월", a1("cloudMonth"), WFMT, "cloud")
arow("", "기본 환경 수용 회원", "명", a1("cloudBaseMembers"), NFMT)
arow("", "초과 회원 1인당 증설 비용", "원/월", a1("cloudPerMember"), WFMT)
arow("AI 상담", "회원 1인당 연 AI 상담 횟수", "회", a1("consultsPerMember"), NFMT, "llm")
arow("", "상담 1회 입력 토큰", "토큰", a1("tokIn"), NFMT)
arow("", "상담 1회 출력 토큰", "토큰", a1("tokOut"), NFMT)
arow("", "LLM 입력 단가(100만 토큰)", "원", a1("priceIn"), D1)
arow("", "LLM 출력 단가(100만 토큰)", "원", a1("priceOut"), WFMT)
arow("기록", "블록체인 기록 건수(2027년)", "건", a1("bcRecords"), NFMT, "bc", once=True)
arow("", "기록 건당 비용", "원/건", a1("bcUnit"), WFMT)
r = section(S2, r, "⑨ 영업·관리비·세무", 9)
arow("요율", "영업비율(매출 대비)", "%", a1("salesRate"), PFMT, "opex")
arow("", "관리비율(매출 대비)", "%", a1("adminRate"), PFMT)
arow("", "표준 요율 대비 적용 비율", "%", a1("opexScale"), PFMT)
arow("세무", "감가상각 내용연수", "년", a1("life"), NFMT, "life")
arow("", "법인세율", "%", a1("tax"), PFMT, "tax")
r = section(S2, r, "⑩ 투자(시스템 구축·고도화·장비)", 9)
AI_LAB = [("checkup", "건강검진"), ("clinic", "진료 안내"), ("care", "재가·돌봄"), ("ins", "보험·치료비"), ("nutri", "건강식단·영양제"),
          ("skin", "피부건강"), ("device", "홈케어 의료기기"), ("sports", "스포츠활동"), ("finance", "경영관리·회계")]
for k, lab in AI_LAB:
    arow("구축" if k == "checkup" else "", f"AI 시스템 구축 — {lab}", "백만원", a1("ai_" + k), MFMT, "capex1" if k == "checkup" else None, once=True)
arow("", "공통 플랫폼 구축", "백만원", a1("aiCommon"), MFMT, None, once=True)
arow("", "보안 인증(ISMS-P 등)", "백만원", a1("isms"), MFMT, None, once=True)
arow("고도화", "시스템 고도화 투자", "백만원", av("capexLater"), MFMT, "capexLater")
arow("장비", "사무 장비(신규 인원 1인당)", "원/인", a1("officePerHead"), WFMT, "office")
arow("", "사무 장비 투자", "백만원", av("capex_office"), MFMT)
put(S2, f"A{r}", "")
put(S2, f"B{r}", "투자 합계", bold=True, indent=1)
put(S2, f"C{r}", "백만원", size=8.5, color=SUBC, h="center")
first_cap = r - (len(AI_LAB) + 5)
for i in range(5):
    col = YC[i]
    put(S2, f"{col}{r}", f"=SUM({col}{first_cap}:{col}{r-1})-{col}{r-2}", fmt=MFMT, bold=True, h="right", exp=av("capexTotal")[i])
put(S2, f"I{r}", "시스템 구축 + 공통 플랫폼 + 보안 인증 + 고도화 + 사무 장비 투자의 합계", size=8.5, color=SUBC)
rule(S2, r, 9, True)
r += 1

# ══════════════════════════ 3. 매출 계획 ══════════════════════════
S3 = wb.create_sheet("3.매출계획", 3)
W3 = [44] + [10] * 12 + [11, 3]
setup(S3, W3)
r = title(S3, 1, "3. 매출 계획", "단위: 백만원", 15)
r = section(S3, r, "① 매출 항목별 산정 방식", 15)
r = header(S3, r, ["매출 항목", "매출 시작", "입금 시점", "산정 방식"], aligns=["left", "center", "center", "left"])
arpu_sum = sum(a1("arpu_" + k) for k, _ in CATS)
METHOD = [("AI 플랫폼 구독", "2027년 3월", "다음 달", "월별 과금 기관 수 × 월 구독료의 월별 합계 · 병원·약국은 2027년 3월, 검진센터는 2028년부터 과금"),
          ("건강검진 연계", "2027년 4월", "다음 달", "건강검진 예약 × (자사 운영 비중 × 건당 3만원 + 제휴사 비중 × 건당 1만원) · 월 예약 = 월말 회원 × 회원당 예약 비율 × 계절 지수"),
          ("예약 서비스(골프·시설)", "2027년 4월", "다음 달", "월말 회원 × 회원당 검진 예약 비율 × 예약 서비스 배수 × 건당 1만원 ÷ 12(계절 지수 미적용, 월별 고르게 배분)"),
          ("헬스메이트센터 사용료", "2027년 4월", "다음 달", "우대 한도 내 공급 × 건당 5만원 + 한도 초과 공급 × 건당 10만원"),
          ("건강커머스", "2027년 7월", "당월", f"월말 회원 × 구매 비율 38% × 1인 연 구매액(4품목 합 {arpu_sum / 1e4:,.0f}만원) ÷ 12 × 당사 채널 70% × 적용률(2027년은 월별 적용률 후 30% 감액)"),
          ("재가·돌봄 파트너 이용료", "2027년 7월", "다음 달", "파트너 센터 수 × 월 이용료 30만원(건당 수수료 없음)")]
for a_, b_, c_, d_ in METHOD:
    put(S3, f"A{r}", a_, indent=1)
    put(S3, f"B{r}", b_, h="center")
    put(S3, f"C{r}", c_, h="center")
    put(S3, f"D{r}", d_, wrap=True)
    S3.merge_cells(start_row=r, start_column=4, end_row=r, end_column=15)
    rule(S3, r, 15)
    r += 1
r = note(S3, r, "※ 연중 회원·제휴 기관 증가에 맞춰 월별로 누적 계산하므로 연간 매출은 연말 규모로 1년을 계산한 금액보다 작음(③ 연말 규모 대비 실제 매출 비율)", 15)
r += 1

r = section(S3, r, "② 연도별 매출", 15)
r = header(S3, r, ["매출 항목", "2027년", "2028년", "2029년", "2030년", "2031년", "5년 합계", "구성비\n(5년 합계)"], aligns=["left"] + ["center"] * 7)
R3 = {}
STREAM3 = [("revSub", "AI 플랫폼 구독"), ("revChk", "건강검진 연계"), ("revResv", "예약 서비스(골프·시설)"), ("revIns", "헬스메이트센터 사용료"),
           ("revP", "건강커머스"), ("revCare", "재가·돌봄 파트너 이용료")]
first3 = r
for k, lab in STREAM3:
    put(S3, f"A{r}", lab, indent=1)
    for i, col in enumerate("BCDEF"):
        put(S3, f"{col}{r}", f"='4.추정손익'!{YC[i]}{R4[k]}", fmt=MFMT, h="right", exp=EXP[f"4.추정손익!{YC[i]}{R4[k]}"])
    put(S3, f"G{r}", f"=SUM(B{r}:F{r})", fmt=MFMT, bold=True, h="right", exp=EXP[f"4.추정손익!I{R4[k]}"])
    R3[k] = r
    rule(S3, r, 8)
    r += 1
put(S3, f"A{r}", "매출액 합계", bold=True)
for i, col in enumerate("BCDEFG"):
    ex = EXP[f"4.추정손익!{YC[i]}{R4['rev']}"] if i < 5 else EXP[f"4.추정손익!I{R4['rev']}"]
    put(S3, f"{col}{r}", f"=SUM({col}{first3}:{col}{r-1})", fmt=MFMT, bold=True, h="right", exp=ex)
rule(S3, r, 8, True)
tot3 = r
for rr in list(R3.values()) + [tot3]:
    put(S3, f"H{rr}", f"=IF($G${tot3}=0,0,G{rr}/$G${tot3})", fmt=PFMT, h="right")
r += 1
put(S3, f"A{r}", "전년 대비 증가율", indent=1)
for i, col in enumerate("CDEF"):
    prev = "BCDE"[i]
    put(S3, f"{col}{r}", f"=IF({prev}{tot3}=0,0,{col}{tot3}/{prev}{tot3}-1)", fmt=PFMT, h="right")
rule(S3, r, 8)
r += 2

r = section(S3, r, "③ 연말 규모 대비 실제 매출 비율 — 연말 회원·기관 수로 1년을 계산한 금액 대비", 15)
r = header(S3, r, ["매출 항목", "2027년", "2028년", "2029년", "2030년", "2031년"], aligns=["left"] + ["center"] * 5)
FULL = {"revSub": "fullSub", "revChk": "fullChk", "revResv": "fullResv", "revIns": "fullIns", "revP": "fullP", "revCare": "fullCare"}
for k, lab in STREAM3:
    full = yv(FULL[k])
    put(S3, f"A{r}", lab + " — 연말 규모 기준 매출", indent=1, color=SUBC)
    for i, col in enumerate("BCDEF"):
        put(S3, f"{col}{r}", full[i], fmt=MFMT, h="right", color=SUBC, exp=full[i])
    rule(S3, r, 6)
    r += 1
    put(S3, f"A{r}", lab + " — 실제 매출 비율", indent=2)
    for i, col in enumerate("BCDEF"):
        put(S3, f"{col}{r}", f"=IF({col}{r-1}=0,0,{col}{R3[k]}/{col}{r-1})", fmt=PFMT, h="right")
    rule(S3, r, 6)
    r += 1
r += 1

S3.row_breaks.append(Break(id=r - 1))
r = section(S3, r, "④ 2027년 월별 매출", 15)
r = header(S3, r, ["매출 항목"] + [f"{m}월" for m in range(1, 13)] + ["합계"], aligns=["left"] + ["center"] * 13)
# 월별자금필요표 2027 열(G..R)
MCOLS27 = [get_column_letter(7 + m) for m in range(12)]
FMAP = [("rSub", "AI 플랫폼 구독", "revSub"), ("rChk", "건강검진 연계", "revChk"), ("rIns", "헬스메이트센터 사용료", "revIns"),
        ("rP", "건강커머스", "revP"), ("rCare", "재가·돌봄 파트너 이용료", "revCare"), ("rOth", "예약 서비스(골프·시설)", "revResv")]
mfirst = r
for fk, lab, k in FMAP:
    put(S3, f"A{r}", lab, indent=1)
    for m in range(12):
        v = F[f"{MCOLS27[m]}{FR[fk]}"].value or 0
        put(S3, f"{get_column_letter(2 + m)}{r}", v, fmt=MFMT, h="right", exp=v)
    put(S3, f"N{r}", f"=SUM(B{r}:M{r})", fmt=MFMT, bold=True, h="right", exp=yv(k)[0])
    rule(S3, r, 14)
    r += 1
put(S3, f"A{r}", "월 매출 합계", bold=True)
for m in range(13):
    col = get_column_letter(2 + m)
    put(S3, f"{col}{r}", f"=SUM({col}{mfirst}:{col}{r-1})", fmt=MFMT, bold=True, h="right", exp=(yv("rev")[0] if m == 12 else None))
rule(S3, r, 14, True)
r += 1
mem_row = next(rr for rr in range(20, 30) if MP[f"B{rr}"].value == "월말 회원(명)")
put(S3, f"A{r}", "월말 회원(명)", indent=1, color=SUBC)
for m in range(12):
    v = MP[f"{get_column_letter(4 + m)}{mem_row}"].value
    put(S3, f"{get_column_letter(2 + m)}{r}", v, fmt=NFMT, h="right", color=SUBC, exp=v)
rule(S3, r, 14)
r += 2

r = section(S3, r, "⑤ 2027년 월별 매출 적용률 · 건강검진 계절 지수 · 2028년 이후 연간 적용률", 15)
r = header(S3, r, ["매출 항목"] + [f"{m}월" for m in range(1, 13)], aligns=["left"] + ["center"] * 12)
FKEYS = [("F_Sub", "AI 플랫폼 구독", "rampSub"), ("F_Chk", "건강검진 연계", "rampChk"), ("F_Resv", "예약 서비스", "rampResv"),
         ("F_Ins", "헬스메이트센터 사용료(월별 배분)", "rampIns"), ("F_P", "건강커머스", "rampP"), ("F_Care", "재가·돌봄", "rampCare")]
MC12 = [get_column_letter(4 + m) for m in range(12)]
for fk, lab, bid in FKEYS:
    put(S3, f"A{r}", lab, indent=1)
    for m in range(12):
        v = A[f"{MC12[m]}{AR[fk]}"].value or 0
        put(S3, f"{get_column_letter(2 + m)}{r}", v, fmt='0%;△0%;"-"', h="right", exp=v)
    rule(S3, r, 13)
    r += 1
put(S3, f"A{r}", "건강검진 월별 계절 지수", indent=1)
for m in range(12):
    v = A[f"{MC12[m]}{AR['season']}"].value
    put(S3, f"{get_column_letter(2 + m)}{r}", v, fmt=PFMT, h="right", exp=v)
rule(S3, r, 13)
r += 1
for fk, lab, bid in FKEYS:
    t_ = f"· {lab}: {BS['ramps'].get(bid, '')}"
    r = note(S3, r, t_, 13, height=15 if len(t_) <= 115 else 26)
r = note(S3, r, f"· 건강검진 월별 계절 지수: {BS['ramps'].get('season', '')}", 13)
r = note(S3, r, "※ 건강커머스 2027년 매출은 위 비율로 계산한 금액에서 30% 감액. 헬스메이트센터 사용료는 2027년 공급 12만건 전량 반영, 위 비율은 월별 배분에만 사용", 15)
r += 1
r = header(S3, r, ["2028년 이후 연간 적용률", "2028년", "2029년", "2030년", "2031년", "근거"], aligns=["left", "center", "center", "center", "center", "left"])
hmerge(S3, r - 1, 6, 13)
for mk, lab in [("M_P", "건강커머스"), ("M_Care", "재가·돌봄")]:
    put(S3, f"A{r}", lab, indent=1)
    for i, col in enumerate("BCDE"):
        v = A[f"{'EFGH'[i]}{AR[mk]}"].value
        put(S3, f"{col}{r}", v, fmt=PFMT, h="right", exp=v)
    put(S3, f"F{r}", BS["ramps"].get("mature_" + ("P" if mk == "M_P" else "Care"), ""), size=8.5, color=SUBC, indent=1)
    S3.merge_cells(start_row=r, start_column=6, end_row=r, end_column=13)
    rule(S3, r, 13)
    r += 1
r = note(S3, r, "※ 그 밖의 매출 항목은 2028년부터 100% 적용", 15)

# ══════════════════════════ 5. 인력 계획 ══════════════════════════
S5 = wb.create_sheet("5.인력계획")
W5 = [36, 10, 12, 12, 12, 12, 12, 52]
setup(S5, W5)
r = title(S5, 1, "5. 인력 계획", "인원: 명(연말) · 금액: 백만원", 8)
SEC = [("checkup", "건강검진"), ("clinic", "진료 안내"), ("care", "재가·돌봄"), ("ins", "보험·치료비(헬스메이트센터 포함)"), ("nutri", "건강식단·영양제"),
       ("skin", "피부건강"), ("device", "홈케어 의료기기"), ("sports", "스포츠활동"), ("finance", "경영관리·회계"), ("cto", "AI·전산 총괄"), ("ceo", "대표이사")]
r = section(S5, r, "① 부문별 인원", 8)
r = header(S5, r, ["부문", "", "2027년", "2028년", "2029년", "2030년", "2031년", "증원 기준(2028년부터)"])
hfirst = r
for k, lab in SEC:
    put(S5, f"A{r}", lab, indent=1)
    vals = hv("head_" + k)
    for i, col in enumerate("CDEFG"):
        put(S5, f"{col}{r}", vals[i], fmt=NFMT, h="right", exp=vals[i])
    put(S5, f"H{r}", BS["driver"].get(k, ""), size=8.5, color=SUBC, indent=1)
    rule(S5, r, 8)
    r += 1
put(S5, f"A{r}", "인원 합계", bold=True)
for i, col in enumerate("CDEFG"):
    put(S5, f"{col}{r}", f"=SUM({col}{hfirst}:{col}{r-1})", fmt=NFMT, bold=True, h="right", exp=hv("headTotal")[i])
put(S5, f"H{r}", "자문위원 4명 별도", size=8.5, color=SUBC, indent=1)
rule(S5, r, 8, True)
htot = r
r += 1
put(S5, f"A{r}", "1인당 매출", indent=1)
for i, col in enumerate("CDEFG"):
    put(S5, f"{col}{r}", f"=IF({col}{htot}=0,0,'4.추정손익'!{YC[i]}{R4['rev']}/{col}{htot})", fmt=MFMT, h="right", exp=hv("rpe")[i])
put(S5, f"H{r}", "매출액 ÷ 연말 인원", size=8.5, color=SUBC, indent=1)
rule(S5, r, 8)
r += 2

r = section(S5, r, "② 부문별 인건비", 8)
r = header(S5, r, ["부문", "", "2027년", "2028년", "2029년", "2030년", "2031년", "비고"])
cfirst = r
for k, lab in SEC:
    put(S5, f"A{r}", lab, indent=1)
    vals = hv("cost_" + k)
    for i, col in enumerate("CDEFG"):
        put(S5, f"{col}{r}", vals[i], fmt=MFMT, h="right", exp=vals[i])
    rule(S5, r, 8)
    r += 1
put(S5, f"A{r}", "자문위원 자문료", indent=1)
for i, col in enumerate("CDEFG"):
    put(S5, f"{col}{r}", hv("adv")[i], fmt=MFMT, h="right", exp=hv("adv")[i])
put(S5, f"H{r}", "4명 × 월 100만원", size=8.5, color=SUBC, indent=1)
rule(S5, r, 8)
r += 1
put(S5, f"A{r}", "인건비 합계", bold=True)
for i, col in enumerate("CDEFG"):
    put(S5, f"{col}{r}", f"=SUM({col}{cfirst}:{col}{r-1})", fmt=MFMT, bold=True, h="right", exp=hv("payTotal")[i])
put(S5, f"H{r}", "4. 추정 손익계산서의 인건비", size=8.5, color=SUBC, indent=1)
rule(S5, r, 8, True)
r += 2

S5.row_breaks.append(Break(id=r - 1))
r = section(S5, r, "③ 2027년 부문별 인원·기본연봉", 8)
r = header(S5, r, ["부문", "인원", "기본연봉(원)", "1인 연 인건비(원)", "", "", "", "기본연봉 산정 근거"], aligns=["left", "center", "center", "center", "center", "center", "center", "left"])
for k, lab in SEC:
    put(S5, f"A{r}", lab, indent=1)
    n_, sal = H[f"D{HR['in_' + k]}"].value, H[f"E{HR['in_' + k]}"].value
    put(S5, f"B{r}", n_, fmt=NFMT, h="right", exp=n_)
    put(S5, f"C{r}", sal, fmt=WFMT, h="right", exp=sal)
    unit = H[f"D{HR['unit_' + k]}"].value
    put(S5, f"D{r}", unit, fmt=WFMT, h="right", exp=unit)
    txt = BS["salary"].get(k, "")
    put(S5, f"E{r}", txt, size=8.5, wrap=True, indent=1)
    S5.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    S5.row_dimensions[r].height = 15 if len(txt) <= 70 else 27
    rule(S5, r, 8)
    r += 1
r += 1
r = section(S5, r, "④ 인건비 산정 기준", 8)
r = header(S5, r, ["항목", "", "2027년", "2028년", "2029년", "2030년", "2031년", "기준"])
put(S5, f"A{r}", "사용자 부담률(4대보험·퇴직급여)", indent=1)
for i, col in enumerate("CDEFG"):
    put(S5, f"{col}{r}", hv("bTotal")[i], fmt=PFMT, h="right", exp=hv("bTotal")[i])
put(S5, f"H{r}", BS["payrule"].get("burden", ""), size=8.5, wrap=True, indent=1)
S5.row_dimensions[r].height = 27
rule(S5, r, 8)
r += 1
hb = lambda k: H[f"D{HR['base_' + k]}"].value
PR = [("1인 연 인건비", "기본연봉 × (1 + 사용자 부담률) + 복리후생 월 25만원 × 12"),
      ("임금 인상률", "연 3.5%(2028년부터 누적 적용)"),
      ("증원 방식", "2028년부터 부문별 담당 지표가 2027년 기준값의 몇 배인지에 비례해 2027년 인원을 늘림(소수점 올림), 대표이사 1명 고정\n"
                f"2027년 기준값: 검진 예약 연 {hb('active') / 1e4:,.0f}만건(연간 규모) · 연말 회원 {hb('me') / 1e4:,.0f}만명 · DB 공급 {hb('insC') / 1e4:,.0f}만건 · "
                f"매출액 {hb('rev') / 1e8:,.0f}억원(상품 매출은 첫해 가동 기간을 1년 기준으로 보정)"),
      ("자문위원", "의료·보험·법률·보험계리 분야 4명, 1인당 월 100만원")]
for a_, b_ in PR:
    put(S5, f"A{r}", a_, indent=1)
    put(S5, f"C{r}", b_, wrap=True)
    S5.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
    if "\n" in b_:
        S5.row_dimensions[r].height = 28
    rule(S5, r, 8)
    r += 1

# ══════════════════════════ 6. 월별 자금 계획 ══════════════════════════
S6 = wb.create_sheet("6.월별자금")
setup(S6, [6, 24, 10.5] + [8.8] * 15)
r = title(S6, 1, "6. 월별 자금 계획 — A 기본 계획", "단위: 백만원 · 2026년 10월 준비 착수, 2027년 1월 서비스 개시 · 필요 자금 확인을 위해 투자금은 2026년 10월 일시 납입으로 계산(실제 분할 납입은 1. 요약 ④)", 18)
MON = [F[f"{get_column_letter(4 + k)}{FR['head']}"].value for k in range(39)]
BLOCKS = [(0, 15, "(가) 2026년 10월~2027년 12월"), (15, 27, "(나) 2028년"), (27, 39, "(다) 2029년")]
ROWSPEC = [
    ("sec", "매출"),
    ("rSub", "AI 플랫폼 구독"), ("rChk", "건강검진 연계"), ("rIns", "헬스메이트센터 사용료"), ("rP", "건강커머스"), ("rCare", "재가·돌봄 파트너 이용료"), ("rOth", "예약 서비스"),
    ("sumr", "매출 합계"),
    ("sec", "입금"),
    ("iSub", "AI 플랫폼 구독"), ("iChk", "건강검진 연계"), ("iIns", "헬스메이트센터 사용료"), ("iP", "건강커머스"), ("iCare", "재가·돌봄 파트너 이용료"), ("iOth", "예약 서비스"),
    ("sumi", "입금 합계"),
    ("sec", "지출"),
    ("oPay", "인건비"), ("oAd", "광고비(개시 전 사전 광고 포함)"), ("oBuild", "시스템 구축(개시 전)"), ("oMedi", "메디에이지 제휴 리포트 구매"),
    ("oIT", "IT 운영·영업·관리비"), ("oCogs", "매출원가"), ("oCust", "회원 포인트 적립·기부금"), ("oRent", "임차료·보증금·채용·설립 비용"),
    ("oCapex", "시스템 고도화·장비 투자"), ("oEtc", "법인세 납부"),
    ("sumo", "지출 합계"),
    ("sec", "자금"),
    ("net", "순현금흐름"), ("cum", "누적 현금(투자 전)"), ("inv", "투자금 납입"), ("bal", "투자 후 현금 잔액"),
]
prev_cum_ref = None
prev_bal_ref = None
for bi, (s_, e_, lab) in enumerate(BLOCKS):
    if bi > 0:
        S6.row_breaks.append(Break(id=r - 1))
    r = section(S6, r, lab, 3 + (e_ - s_))
    hdr = ["구분", "항목", "기간 합계"] + [f"{str(MON[k])[2:4]}.{str(MON[k])[5:7]}" for k in range(s_, e_)]
    r = header(S6, r, hdr, aligns=["left", "left"] + ["center"] * (len(hdr) - 2))
    RB = {}
    grp_first = None
    for key, label in ROWSPEC:
        n = e_ - s_
        lastc = get_column_letter(3 + n)
        if key == "sec":
            put(S6, f"A{r}", label, bold=True, color=NAVY)
            grp_first = r + 1
            r += 1
            continue
        put(S6, f"B{r}", label, indent=0 if key.startswith("sum") or key in ("net", "cum", "bal") else 1,
            bold=key.startswith("sum") or key in ("net", "cum", "bal"))
        total = key.startswith("sum") or key in ("net", "cum", "bal")
        for j, k in enumerate(range(s_, e_)):
            col = get_column_letter(4 + j)
            src = F[f"{get_column_letter(4 + k)}{FR[{'sumr': 'rTot', 'sumi': 'iTot', 'sumo': 'oTot'}.get(key, key)]}"].value or 0
            if key.startswith("sum"):
                f_ = f"=SUM({col}{grp_first}:{col}{r-1})"
            elif key == "net":
                f_ = f"={col}{RB['sumi']}-{col}{RB['sumo']}"
            elif key == "cum":
                if j == 0:
                    f_ = f"={col}{RB['net']}" if prev_cum_ref is None else f"={prev_cum_ref}+{col}{RB['net']}"
                else:
                    f_ = f"={get_column_letter(3 + j)}{r}+{col}{RB['net']}"
            elif key == "bal":
                f_ = f"={col}{RB['cum']}+SUM($D${RB['inv']}:{col}{RB['inv']})" + (f"+{prev_inv_total}" if (bi > 0 and prev_inv_total) else "")
            else:
                f_ = src
            put(S6, f"{col}{r}", f_, fmt=MFMT, h="right", bold=total, exp=src)
        if key not in ("cum", "bal"):
            srcsum = sum((F[f"{get_column_letter(4 + k)}{FR[{'sumr': 'rTot', 'sumi': 'iTot', 'sumo': 'oTot'}.get(key, key)]}"].value or 0) for k in range(s_, e_))
            put(S6, f"C{r}", f"=SUM(D{r}:{lastc}{r})", fmt=MFMT, h="right", bold=True, exp=srcsum)
        rule(S6, r, 3 + n, total)
        RB[key] = r
        r += 1
    n = e_ - s_
    prev_cum_ref = f"{get_column_letter(3 + n)}{RB['cum']}"
    prev_inv_total = (f"{('C' + str(RB['inv']))}" if bi == 0 else f"{prev_inv_total}+C{RB['inv']}")
    r += 1
r = note(S6, r, "※ 누적 현금(투자 전) 최저 달의 부족액에 6개월분 고정비를 더해 투자 요청액 산정(7. 투자금 산정). 법인세는 전년도분을 3월에 납부", 18)

# ══════════════════════════ 7. 투자금 산정 ══════════════════════════
S7 = wb.create_sheet("7.투자금산정")
W7 = [36, 16, 16, 16, 66]
setup(S7, W7)
r = title(S7, 1, "7. 투자금 산정", "단위: 백만원", 5)
r = section(S7, r, "① 산정 원칙", 5)
PRINC = ["월별로 매출 입금과 지출을 쌓아 누적 현금이 가장 낮아지는 달의 부족액을 구함(투자금 없이 운영한다고 가정)",
         "매출이 계획보다 늦어질 때를 대비해 2027년 월평균 고정비(인건비·IT 운영비·관리비)의 6개월분을 안전 버퍼로 더함",
         "필요 총자금(최저점 부족액 + 안전 버퍼)을 10억원 단위로 올린 금액을 투자 요청액으로 정함",
         "보유 현금은 차감하지 않고, 차입금은 없는 것으로 계산"]
for i_, t_ in enumerate(PRINC):
    r = note(S7, r, f"{i_ + 1}) {t_}", 5, size=9, color=INK)
r += 1
r = section(S7, r, "② 계획별 전제", 5)
r = header(S7, r, ["구분", "A 기본 계획", "B 보수 계획", "C 지연 계획", "설명"], aligns=["left", "center", "center", "center", "left"])
pre = cv("pre")
PREM = [("준비 착수", "2026년 10월", "2026년 10월", "2026년 10월", "세 계획 모두 같은 달부터 비용 발생"),
        ("서비스 개시", "2027년 1월", "2027년 3월", "2027년 3월", "B·C는 개시가 2개월 늦어지고 매출 일정도 함께 늦어짐"),
        ("헬스메이트센터 사용료 입금", "다음 달", "2개월 후", "2개월 후", "보험사 정산이 늦어지는 경우 가정"),
        ("시스템 고도화 투자 집행", "월별 균등", "분기 초 선집행", "분기 초 선집행", ""),
        ("헬스메이트센터 DB 공급", "계획대로", "계획대로", "개시 후 12개월 없음", "법률 검토·시범 운영 판정이 늦어지는 경우")]
for row_ in PREM:
    for i, v in enumerate(row_):
        put(S7, f"{'ABCDE'[i]}{r}", v, h="left" if i in (0, 4) else "center", indent=1 if i == 0 else 0, size=9 if i < 4 else 8.5, color=INK if i < 4 else SUBC)
    rule(S7, r, 5)
    r += 1
r += 1
r = section(S7, r, "③ 계획별 소요자금", 5)
r = header(S7, r, ["구분", "A 기본 계획", "B 보수 계획", "C 지연 계획", "산정 방식"], aligns=["left", "center", "center", "center", "left"])
R7 = {}
for key, lab, vals, fmt, rem in [("low", "누적 현금 최저점(투자 전)", low, MFMT, "6. 월별 자금 계획의 누적 현금 최저값"),
                                  ("lowcal", "최저점 발생 월", lowcal, None, ""),
                                  ("fixed", "2027년 월평균 고정비", fixed, MFMT, "인건비·IT 운영비(유지보수·데이터·보안·클라우드 기본)·관리비"),
                                  ("buf", "안전 버퍼", None, MFMT, "월평균 고정비 × 6개월"),
                                  ("need", "필요 총자금", None, MFMT, "최저점 부족액 + 안전 버퍼"),
                                  ("req", "투자 요청액", None, MFMT, "필요 총자금을 10억원 단위로 올림")]:
    put(S7, f"A{r}", lab, indent=1, bold=key in ("need", "req"))
    for i, col in enumerate("BCD"):
        if key == "buf":
            put(S7, f"{col}{r}", f"=ROUND({col}{R7['fixed']}*6,0)", fmt=fmt, h="right", exp=buf[i])
        elif key == "need":
            put(S7, f"{col}{r}", f"=-{col}{R7['low']}+{col}{R7['buf']}", fmt=fmt, h="right", bold=True, exp=need[i])
        elif key == "req":
            put(S7, f"{col}{r}", f"=ROUNDUP({col}{R7['need']}/1000000000,0)*1000000000", fmt=fmt, h="right", bold=True, fill=KEYF, exp=req[i])
        else:
            put(S7, f"{col}{r}", vals[i], fmt=fmt, h="right" if fmt else "center", exp=(vals[i] if fmt else None))
    put(S7, f"E{r}", rem, size=8.5, color=SUBC, indent=1)
    rule(S7, r, 5, key in ("need", "req"))
    R7[key] = r
    r += 1
r += 1
S7.row_breaks.append(Break(id=r - 1))
r = section(S7, r, "④ 현금 흐름 가정", 5)
r = header(S7, r, ["항목", "적용 기준", "", "", "근거"], aligns=["left", "left", "left", "left", "left"])
for a_, b_, bid in [("매출 입금 시점", "건강커머스 당월(카드·간편결제 정산), 그 밖의 매출 다음 달", "lag"),
                    ("매출원가 지급", "발생한 달에 지급(공급사 외상 기간 미반영)", ""),
                    ("준비기간 비용(2026년 10~12월)", "인건비 60%·80%·100%, 사전 광고 30%·60%·100%, IT 운영 11월부터, 시스템 구축비 30%·40%·30%", "pre"),
                    ("임차료·보증금", "1인당 월 54만원(관리비 포함), 보증금 월 임차료 10개월분", "rent"),
                    ("채용 비용", "신규 채용 인원 1인 연 인건비의 5%", "hire"),
                    ("설립·거래 비용", "실사·법무·등기·사무실 준비 1.5억원(2026년 10월)", "oneoff"),
                    ("법인세 납부", "전년도분을 다음 해 3월에 납부, 이월결손금 공제 후 22%", "tax"),
                    ("안전 버퍼", "2027년 월평균 고정비 6개월분", "buffer")]:
    put(S7, f"A{r}", a_, indent=1)
    put(S7, f"B{r}", b_, wrap=True)
    S7.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    txt = basis(bid) if bid else ""
    put(S7, f"E{r}", txt, size=8.5, color=SUBC, wrap=True)
    S7.row_dimensions[r].height = 15 if max(len(b_), len(txt)) <= 45 else 27
    rule(S7, r, 5)
    r += 1
r += 1
r = section(S7, r, "⑤ 법인세 계산 — A 기본 계획", 5)
r = header(S7, r, ["연도", "세전이익", "이월결손금 공제", "법인세", "비고"], aligns=["left", "center", "center", "center", "left"])
for i in range(5):
    yr = 2027 + i
    put(S7, f"A{r}", f"{yr}년", indent=1)
    put(S7, f"B{r}", f"='4.추정손익'!{YC[i]}{R4['pbt']}", fmt=MFMT, h="right", exp=yv("pbt")[i])
    put(S7, f"C{r}", f"='4.추정손익'!{YC[i]}{R4['nol']}", fmt=MFMT, h="right", exp=yv("nolOpen")[i])
    put(S7, f"D{r}", f"=MAX(0,B{r}-C{r})*22%", fmt=MFMT, h="right", exp=yv("tax")[i])
    put(S7, f"E{r}", f"{yr + 1}년 3월 납부" + (f" · 2026년 준비기간 비용 {eok(nol27)}(구축비·보증금 제외)을 이월결손금으로 공제" if i == 0 else ""), size=8.5, color=SUBC, indent=1)
    rule(S7, r, 5)
    r += 1

# ══════════════════════════ 8. 단가 산정 근거 ══════════════════════════
S8 = wb.create_sheet("8.산정근거")
W8 = [18, 26, 30, 58, 42]
setup(S8, W8, title_rows="4:4")
r = title(S8, 1, "8. 단가 산정 근거", "조사 기준일 2026년 9월 · 금액은 부가가치세 별도", 5)
r = header(S8, 4, ["구분", "항목", "적용값", "산정 근거", "출처"], aligns=["left"] * 5)
r = 5
for sec_name, rows in BS["evidence"]:
    r = section(S8, r, sec_name, 5)
    last_grp = None
    for row_ in rows:
        grp = row_["group"] if row_["group"] != last_grp else ""
        last_grp = row_["group"]
        put(S8, f"A{r}", grp, bold=True, color=SUBC, wrap=True)
        put(S8, f"B{r}", row_["item"], wrap=True)
        put(S8, f"C{r}", row_["applied"], wrap=True)
        put(S8, f"D{r}", row_["basis"], size=8.5, wrap=True)
        srcs = [s for s in row_["sources"] if s.get("name")]
        put(S8, f"E{r}", "\n".join(s["name"] for s in srcs), size=8, color=SUBC, wrap=True)
        if srcs and srcs[0].get("url"):
            S8[f"E{r}"].hyperlink = srcs[0]["url"]
        import math
        lines_ = max(math.ceil(len(row_["basis"]) / 46), sum(max(1, math.ceil(len(x["name"]) / 43)) for x in srcs) if srcs else 1,
                     math.ceil(len(row_["item"]) / 16), math.ceil(len(row_["applied"]) / 21), 1)
        S8.row_dimensions[r].height = min(11.5 * lines_ + 5, 110)
        rule(S8, r, 5)
        r += 1

# 시트 순서 정리
order = ["표지", "1.요약", "2.주요가정", "3.매출계획", "4.추정손익", "5.인력계획", "6.월별자금", "7.투자금산정", "8.산정근거"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT)
io.open(OUT + ".expect.json", "w", encoding="utf-8").write(json.dumps({k: v for k, v in EXP.items() if v is not None}, ensure_ascii=False))
print("saved", OUT, "· 기대값", len([v for v in EXP.values() if v is not None]))
