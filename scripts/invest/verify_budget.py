# -*- coding: utf-8 -*-
"""투자금 양식 검증 — 엑셀(재계산 후 저장본) 값을 독립 정답지(oracle.py · oracle3.py)와 대조한다.
사용: python scripts/invest/verify_budget.py docs/invest/하이젠케어_투자금산정_예산양식_vX.Y.xlsx
통과 기준: 수식 오류 0 · 값 없음 0 · 연간 계정 차이 0원 · 월별→연간 차이 0원 · A/B/C 저점·필요·요청·시기·런웨이 일치"""
import sys, json, os
sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as CL
import oracle as O, oracle3 as O3

X = sys.argv[1]
mp = json.load(open(X + ".map.json", encoding="utf-8"))
wv = load_workbook(X, data_only=True); wf = load_workbook(X)
n = e = z = 0
for ws in wf.worksheets:
    v = wv[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n += 1; val = v[c.coordinate].value
                if val is None: z += 1
                elif isinstance(val, str) and val.startswith("#"): e += 1; print("  오류", ws.title, c.coordinate, val)
print("수식 %d개 · 오류 %d · 값 없음 %d" % (n, e, z))
YR, CR, MR, TR = mp["YR"], mp["CR"], mp["MR"], mp["TR"]
Y = wv["연간손익"]; C = wv["현금·투자금"]; M = wv["월별예산"]
adj = O.annual(O3.P0)
pairs = [("insCases", "insC"), ("hmCap", "hmCap"), ("hmPrice", "hmPrice"), ("hmDiscN", "hmDisc"), ("hmFullN", "hmFull"),
         ("hmBenefit", "hmBenefit"), ("revIns", "revIns"), ("revChk", "revChk"), ("revSvc", "revSvc"), ("revSub", "revSub"),
         ("rev", "rev"), ("cogs", "cogs"), ("gross", "gross"), ("sga", "sga"), ("ebit", "ebit")]
for a in adj:
    a["revIns_d"] = a["hmDisc"] * a["hmPrice"]; a["revIns_f"] = a["hmFull"] * a["hmMarket"]
pairs += [("revIns_d", "revIns_d"), ("revIns_f", "revIns_f"), ("imRev", "imRev")]
if "startF" in O3.P0:                              # v3.0 매출 개시 일정 · 월별 누적 · 재가·돌봄 · 연 환산
    for a in adj:
        a["insCasesFull"] = a["insCFull"]; a["activeAct"] = a["active"]; a["fullRev"] = a["revFull"]; a["fullP"] = a["revPFull"]
        a["fullChk"] = a["revChkFull"]; a["fullSub"] = a["revSubFull"]; a["fullIns"] = a["revInsFull"]; a["fullCare"] = a["revCareFull"]
        a["revP_"] = a["revP"]
    pairs += [(k, k) for k in ("revP", "revResv", "revCare", "careCost", "resv", "careN", "careCtrFull", "activeAct", "insCasesFull",
                               "fullRev", "fullP", "fullChk", "fullSub", "fullIns", "fullCare")]
if "cost2" in O3.P0:                               # 판관비·CAPEX 근거 모델(v2.0)
    for a in adj:
        a["capexMedi"] = a["medi"]; a["capexBuild1"] = a["build1"]
    pairs += [(k, k) for k in ("headTotal", "pay", "mk1", "media", "creative", "cardAd", "kit", "sticker", "qrfee", "mktSum",
                               "itMaint", "itData", "itSec", "cloudBase", "cloudVar", "llm", "bc", "itOpex", "capex", "capexMedi", "capexBuild1", "depr")]
worst = max((abs(Y[f"{'DEFGH'[i]}{YR[xk]}"].value - adj[i][ok]), xk, i + 1) for xk, ok in pairs if xk in YR for i in range(5))
if "startF" in O3.P0:                              # 매출 반영률(엑셀 SUM vs 정답지) — 부동소수 차이만 허용
    import adjust as _A
    effd = max(abs(Y[f"{'DEFGH'[i]}{YR['eff_'+k]}"].value - adj[i]["eff"][k]) for k, _, _ in _A.STREAMS for i in range(5))
    print("매출 반영률 7줄×5년 최대 차이: %.2e" % effd)
    worst = max(worst, (1 if effd > 1e-9 else 0, "eff", 0))
cumb = 0.0; bx = 0.0
for i in range(5):
    cumb += adj[i]["hmBenefit"]
    bx = max(bx, abs(Y[f"{'DEFGH'[i]}{YR['hmBenefitX']}"].value - cumb / O3.P0["hmInvest"]))
print("투자금 대비 할인 배수 최대 차이: %.2e" % bx)
print("연간 계정 엑셀 vs 정답지 최대 차이(원, 계정, 연차):", worst)
wm = max((abs(M[f"{CL(4+65+i)}{r}"].value or 0) for r in MR.values() for i in range(5) if isinstance(M[f"{CL(4+65+i)}{r}"].value, (int, float))), default=0)
print("월별→연간 최대 차이: %.6f원" % wm)
res = [O3.run(O3.P0, L) for L in (O3.LA, O3.LB, O3.LC)]
ok = all(abs(C[f"{c}{CR['s_'+k]}"].value - r[key]) <= 1 for k, key in [("low", "low"), ("need", "need"), ("req", "req")] for c, r in zip("CDE", res))
ok &= [C[f"{c}{CR['s_lowAt']}"].value for c in "CDE"] == [r["low_at"] for r in res]
ok &= [C[f"{c}{CR['s_runway']}"].value for c in "CDE"] == [r["runway"] for r in res]
# 월별 누적 현금 전 구간 대조(블록별)
gmax = 0.0
for R, r in zip((mp["RA"], mp["RB"], mp["RC"]), res):
    for j, s in enumerate(r["series"]):
        gmax = max(gmax, abs(C.cell(row=R["cum"], column=3 + j).value - s[4]))
ok &= gmax <= 1
imax = 0.0
for R, r in zip((mp["RA"], mp["RB"], mp["RC"]), res):
    for t in range(60):
        imax = max(imax, abs(C.cell(row=R["insRec"], column=3 + 24 + t).value - r["insRecM"][t]))
ok &= imax <= 1 and bx < 1e-9
print("사용료 현금 인식(insRec) 60개월×3 최대 차이: %.4f원" % imax)
print("요청액:", [C[f"{c}{CR['s_req']}"].value / 1e8 for c in "CDE"], "· 필요:", [round(C[f"{c}{CR['s_need']}"].value / 1e8, 1) for c in "CDE"],
      "· 저점:", [(round(C[f"{c}{CR['s_low']}"].value / 1e8, 1), C[f"{c}{CR['s_lowAt']}"].value) for c in "CDE"],
      "· 런웨이:", [C[f"{c}{CR['s_runway']}"].value for c in "CDE"])
print("월별 누적 현금 84개월×3 최대 차이: %.4f원" % gmax, "· 정답지 일치:", "✓" if ok else "✗")
for i in range(7 if "startF" in O3.P0 else 6):
    print("점검", C[f"B{CR['warnStart']+i}"].value)
if "taxes" in res[0]:                              # 법인세 이월결손금 — 연간손익(계획 = A) vs 정답지 A 연도별
    tx = max(abs(Y[f"{'DEFGH'[i]}{YR['tax']}"].value - res[0]["taxes"][i]) for i in range(5))
    txs = max(abs(C[f"{c}{CR['s_tax'+str(k+1)]}"].value - r["taxes"][k]) for c, r in zip("CDE", res) for k in range(4))
    print("법인세(이월결손금) 연간손익 5년 최대 차이: %.4f원 · 현금 요약 A·B·C 4년 최대 차이: %.4f원" % (tx, txs))
    ok &= tx <= 1 and txs <= 1
# 월별 자금필요표(기준 선택 = A) — 정답지 A 시리즈와 36개월 대조
if "FR" in mp and "월별자금필요표" in wv.sheetnames:
    FT = wv["월별자금필요표"]; FR = mp["FR"]; rA = res[0]; preA = O3.LA["pre"]
    NM = 39 if "startF" in O3.P0 else 36
    MMC = [CL(4 + k) for k in range(NM)]
    tmax = 0.0
    for k in range(NM):
        lab, idx, inflow, outflow, cumv = rA["series"][24 - preA + k]
        tmax = max(tmax, abs(FT[f"{MMC[k]}{FR['iTot']}"].value - inflow), abs(FT[f"{MMC[k]}{FR['oTot']}"].value - outflow),
                   abs(FT[f"{MMC[k]}{FR['cum']}"].value - cumv), abs(FT[f"{MMC[k]}{FR['chk']}"].value or 0))
    cal0 = FT[f"D{FR['head']}"].value
    ok_t = FT["C4"].value == 1 and tmax <= 1 and cal0 == ("2026-10" if NM == 39 else "2026-11")
    ok &= ok_t
    print("월별 자금필요표(A) %d개월 유입·유출·누적 최대 차이: %.4f원 · 첫 달 %s · 오픈 첫 달 %s · %s" % (NM, tmax, cal0, FT[f"C{FR['minBal']+1}"].value, "✓" if ok_t else "✗"))
sys.exit(0 if (e == 0 and z == 0 and worst[0] == 0 and wm < 1e-3 and ok) else 1)
